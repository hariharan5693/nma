"""One-off template updater for the nma global template.

Adds, without disturbing the existing field-definition sheets:
  1. A new meta sheet  SHEET_DESCRIPTIONS  — one row per DOMAIN sheet with a
     file/sheet-level BUSINESS_DESCRIPTION + BUSINESS_FACTS (adapted from the
     India catalog where a business-equivalent exists), plus GEOGRAPHY_LEVELS
     and MANUFACTURER_SCOPE columns (points 1 + 2).
  2. A GEOGRAPHY_LEVEL field row into DIM_GEOGRAPHY (point 2 — geography grain).
  3. Manufacturer / import-origin standard fields into the fact sheets that need
     them: Secondary Volume, Financial Shipments, Industry Segmentation, Nielsen
     (point 2 — manufacturer + imports).
  4. A BUSINESS_DESCRIPTION_LONG column into DIM_FILE and DIM_SHEET (point 1).

Idempotent: re-running will not duplicate rows/columns/sheets it already added.
Run from nma/ with:  .venv\\Scripts\\python _update_template.py
"""

from __future__ import annotations

import openpyxl

TPL = "templates/global_template.xlsx"

# Field-sheet header is fixed: A..G
HDR = [
    "STANDARD_FIELD_NAME", "BUSINESS_DESCRIPTION", "FIELD_TYPE", "FLAG",
    "INDIA_FIELD_NAME", "THAILAND_FIELD_NAME", "VIETNAM_FIELD_NAME",
]

# ──────────────────────────────────────────────────────────────────────────
# 1 + 2. Per-domain file/sheet business descriptions, geography grain, mfg scope.
#    Descriptions adapted from India catalog business meaning; written generic
#    (market-agnostic) so they describe the STANDARD domain, not India specifics.
# ──────────────────────────────────────────────────────────────────────────
SHEET_DESCRIPTIONS: list[dict[str, str]] = [
    {
        "domain": "PMI P&L Database",
        "desc": "Full profit-and-loss database across data versions (Actuals vs Actuals, "
                "vs Original Budget, vs Rolling Forecast). Carries current and reference "
                "volumes/amounts plus a variance-bridge decomposition (Volume, Mix, Price, "
                "FX, Real, Other) that explains why the current period differs from a "
                "reference scenario at any P&L line, with a Mapping sheet that assigns each "
                "account line to NR -> MC -> GP -> PC waterfall levels.",
        "facts": "Financial file (not volume sell-in or sell-out). Enables revenue, margin, "
                 "cost and profitability analysis by brand family/differentiator. Comparisons: "
                 "AC vs AC (vs previous year), AC vs OB (vs annual plan), AC vs RF (vs latest "
                 "forecast). Real = Volume+Mix+Price+ex-Mnf+CC409+SVC+FME+Other; Total = FX+Real.",
        "geo": "Market Group > Market",
        "mfg": "PMI own financials (single manufacturer).",
    },
    {
        "domain": "PMI Nielsen Dashboard",
        "desc": "Third-party retail-audit / scandata market data (Nielsen or equivalent "
                "retail panel: 7-Eleven, Tops, Key-Account scandata) giving sell-out sales "
                "value/volume, market share, store counts and distribution metrics across "
                "manufacturers, brands, variants, sizes, strengths and price segments, by "
                "market and period. ANY retail-panel or scandata sell-out-by-store file "
                "belongs here.",
        "facts": "Competitive retail-audit data covering PMI AND all competitors. Measures "
                 "true market demand, brand pull and share. Additive facts: store count, "
                 "sales volume, sales value, stocks. Share/distribution %s are ratios and "
                 "must not be summed. This is sell-out (consumer off-take), NOT shipment.",
        "geo": "Market / Zone / City",
        "mfg": "All manufacturers (TMO_MANUFACTURER dimension).",
    },
    {
        "domain": "PMI Secondary Volume",
        "desc": "In-market sales (IMS) / secondary sell-out volume: volume sold from the "
                "distributor/wholesaler to retail outlets. Granular by brand, SKU, variant, "
                "state/province, city and channel (Dealer/Hawker/Modern/Retail/Key-Account), "
                "captured monthly. Represents true market demand and consumption.",
        "facts": "Sell-OUT (distributor -> retailer), NOT sell-in. Distinct from Financial "
                 "Shipments which is PMI -> distributor. A file with distributor/account/"
                 "region/channel dimensions and per-channel volume (e.g. CN/INV net units, "
                 "Dealer/Hawker/Modern) is Secondary Volume even if its filename says "
                 "'Shipment'. Links to Trade Offers and SFA/UBO by state/city/SKU.",
        "geo": "Region > State/Province > City/Town",
        "mfg": "PMI own sell-out (single manufacturer).",
    },
    {
        "domain": "PMI Financial Shipments",
        "desc": "Primary shipment volume: volume sold FROM PMI TO its primary distributor "
                "(sell-in). Historic plus forecast month-on-month shipment volumes per SKU/"
                "variant, used for forecast-accuracy (SFA) and bias, and linked to capacity "
                "and inventory by SKU format.",
        "facts": "Sell-IN (PMI -> distributor), NOT sell-out. A SKU-level, forecastable "
                 "monthly volume file with product codes and brand totals (e.g. IBP5 cycle "
                 "files) is Financial Shipments. A distributor-to-retail channel file is NOT "
                 "- that is Secondary Volume. Shipment moves lead secondary volume by 1-2 "
                 "months.",
        "geo": "National (by SKU/variant); optional production/destination country for imports.",
        "mfg": "PMI own shipments; production vs destination country distinguishes imports.",
    },
    {
        "domain": "PMI OB and RF Volume",
        "desc": "Original Budget (OB) and Rolling Forecast (RF) volume plans per product, "
                "laid out as monthly/quarterly planned volumes by data version.",
        "facts": "Planning file. Compare actual volume (Secondary Volume / Financial "
                 "Shipments) against OB (annual plan) and RF (latest forecast). OB fixed for "
                 "the year; RF revised quarterly (RF3..RF11).",
        "geo": "Market (plan level).",
        "mfg": "PMI own plan (single manufacturer).",
    },
    {
        "domain": "PMI India Industry Segmentation",
        "desc": "Total-industry cigarette volume across all trademark owners (TMOs), "
                "structured by product dimensions (brand, length/price/format/flavor segment) "
                "at the EBROM/SKU level with monthly and YTD volumes. Primary source for "
                "Share of Market and Share of Segment. Industry sell-in / shipment basis, not "
                "consumer off-take.",
        "facts": "Industry-wide (all TMOs), not PMI-only. Use for SOM/SOS, competitive "
                 "benchmarking, pricing architecture and substitution. GSPR IMS / industry "
                 "volume feeds (by TMO + EBROM) belong here. Note shipment basis for "
                 "like-to-like vs secondary sales.",
        "geo": "National / market total (by TMO and segment).",
        "mfg": "All manufacturers / TMOs (TMO_NAME, TMO_CODE dimensions).",
    },
    {
        "domain": "PMI vs Competition Mapping",
        "desc": "Brand/variant mapping tables: map own (PMI) SKUs/variants to the closest "
                "competitor SKUs/variants across length, price, format and flavor segments; "
                "often one tab per TMO (PMI, JTI, BAT, etc.), plus TMO -> brand -> sub-family "
                "-> variant hierarchy tables.",
        "facts": "Reference/mapping file (names, not volumes). Used to align own vs "
                 "competitor products for share and mix-movement analysis and to resolve "
                 "which variant falls under which length/price segment. Per-competitor tabs "
                 "(PMI, JTI, TTM, Leadon, Habanos, Others) all map to this domain.",
        "geo": "N/A (product taxonomy).",
        "mfg": "PMI and all competitor TMOs.",
    },
    {
        "domain": "PMI Capacity Glidepath",
        "desc": "Manufacturing capacity forecast in volume per month by SKU format over a "
                "rolling horizon; tracks capacity planning and ramp-up.",
        "facts": "Benchmark capacity vs Financial Shipments and Inventory at SKU-FORMAT "
                 "level (shared production line, not per-SKU). Shipments >= ~85% of capacity "
                 "signals a constraint.",
        "geo": "Plant / SKU format.",
        "mfg": "PMI own capacity.",
    },
    {
        "domain": "PMI Inventory Details",
        "desc": "Month-opening inventory volume held by PMI and/or by the primary "
                "distributor, by SKU / variant / length-format / SKU format, month by month.",
        "facts": "Month-OPENING inventory (not average/closing). Holder entity distinguishes "
                 "PMI vs distributor stock. Combine with shipments for Days of Inventory "
                 "(DOI) and with capacity by SKU format.",
        "geo": "Holder entity / warehouse.",
        "mfg": "PMI own inventory (PMI and distributor holders).",
    },
    {
        "domain": "PMI SFA UBO",
        "desc": "Unique Billed Outlets (UBO): count of retail outlets selling PMI brands/"
                "SKUs/variants across states and cities, by outlet type, month by month.",
        "facts": "Route-to-market strength indicator. Correlate UBO changes with Secondary "
                 "Volume and Trade Offers by state/city/SKU/variant.",
        "geo": "State > City > Town (outlet level).",
        "mfg": "PMI own distribution footprint.",
    },
    {
        "domain": "PMI Trade Offers",
        "desc": "Trade offer / promotion records by region, branch, state, city cluster and "
                "SKU: offer type, period, beneficiary, applicable channel and distribution "
                "node, plus an SKU master.",
        "facts": "Promotion deployment reference. Link to Secondary Volume and SFA/UBO to "
                 "measure promotion impact on sell-out and outlet reach.",
        "geo": "Region > Branch > State > City cluster.",
        "mfg": "PMI own trade programs.",
    },
    {
        "domain": "PMI Key Indicators",
        "desc": "Headline KPI scorecard (SFA, Bias, Customer Fill Rate, Productivity, LOGD, "
                "Ageing Inventory, Innovation Index, KS SOS, Brand Equity, distribution "
                "reach, etc.) by month, with FY target and YTD benchmarks.",
        "facts": "Single source for headline KPIs. Thresholds: SFA < 90% or Bias > 5% or "
                 "Fill Rate < 99% flagged red. FY Target and FY YTD columns benchmark "
                 "performance.",
        "geo": "Market total.",
        "mfg": "PMI own KPIs.",
    },
    {
        "domain": "PMI NCPT BAM Funnel",
        "desc": "Brand-awareness-to-buying (BAM) funnel metrics from consumer research: "
                "Total Awareness, Tried P12M, Bought P12M, Penetration and Main Brand usage, "
                "by segment or brand, geography set, research wave and age cohort.",
        "facts": "Consumer survey funnel (percentages; base counts are absolute). Assesses "
                 "brand health via awareness-to-purchase conversion across cities and cohorts.",
        "geo": "Focus-city sets (e.g. 5 FC / 11 FC).",
        "mfg": "PMI and competitor brands (survey).",
    },
    {
        "domain": "REGULATORY_COMPLIANCE",
        "desc": "Regulatory submissions, approvals, licenses, advertising/zone restrictions "
                "and tax rates per product and period for the market.",
        "facts": "Compliance reference. Tracks authority, submission/approval status, license "
                 "validity, excise/VAT rates and violations.",
        "geo": "Country / restricted zones.",
        "mfg": "PMI regulatory position.",
    },
    {
        "domain": "MARKET_MATURITY",
        "desc": "Market transformation-stage metrics: smoke-free launch year, years since "
                "launch, maturity stage, categories available, and smoke-free revenue/volume "
                "share with threshold flags.",
        "facts": "Strategic transformation tracking per market vs 2030 targets.",
        "geo": "Market / country.",
        "mfg": "PMI transformation position.",
    },
    {
        "domain": "SMOKE_FREE_USERS",
        "desc": "Smoke-free product user / switcher tracking: total smoke-free users by "
                "category (HTU/E-Vapor/Oral), adult smokers who switched, penetration, "
                "acquisition, churn and retention.",
        "facts": "Strategic transformation user base tracking.",
        "geo": "Market / optional geography.",
        "mfg": "PMI smoke-free users.",
    },
    {
        "domain": "STRATEGIC_TRANSFORMATION_KPIs",
        "desc": "PMI 2030 transformation financial ratios and volumes: smoke-free vs total "
                "net revenues, R&D, commercial spend, capex ratios, markets with SF products "
                "and smoke-free shipment volumes.",
        "facts": "Corporate 2030 transformation scorecard vs targets.",
        "geo": "Country / corporate reporting level.",
        "mfg": "PMI corporate.",
    },
]

# ──────────────────────────────────────────────────────────────────────────
# 3. Manufacturer / import-origin fields to ADD into fact sheets (if absent).
#    (standard_field_name, business_description, field_type, flag, IN, TH, VN)
# ──────────────────────────────────────────────────────────────────────────
MFG_IMPORT_FIELDS: dict[str, list[tuple]] = {
    "PMI Secondary Volume": [
        ("MANUFACTURER", "Manufacturer / trademark owner of the product (PMI vs competitor). "
         "Defaults to PMI for own sell-out.", "ATTR", "C", "PMI", "PMI", "PMI"),
        ("SOURCE_ORIGIN", "Domestic vs imported origin of the product (for markets where "
         "some volume is imported).", "ATTR", "O", "Domestic", "Domestic/Imported", "Domestic"),
        ("IMPORT_FLAG", "Flag indicating the product/volume is imported rather than "
         "locally produced.", "ATTR", "O", "", "", ""),
    ],
    "PMI Financial Shipments": [
        ("MANUFACTURER", "Manufacturer / trademark owner (PMI for own shipments).",
         "ATTR", "C", "PMI", "PMI", "PMI"),
        ("IMPORT_FLAG", "Imported vs domestically produced shipment (production country "
         "differs from destination country).", "ATTR", "O", "", "", ""),
    ],
    "PMI India Industry Segmentation": [
        ("SOURCE_ORIGIN", "Domestic vs imported origin of the industry volume (where "
         "tracked).", "ATTR", "O", "", "", ""),
        ("IMPORT_FLAG", "Flag indicating imported industry volume.", "ATTR", "O", "", "", ""),
    ],
    "PMI Nielsen Dashboard": [
        ("SOURCE_ORIGIN", "Domestic vs imported origin of the product (where the retail "
         "panel distinguishes it).", "ATTR", "O", "", "", ""),
    ],
}


def _header_row(ws) -> list[str]:
    return [(str(c.value).strip().upper() if c.value is not None else "")
            for c in ws[1]]


def update_sheet_descriptions(wb) -> None:
    name = "SHEET_DESCRIPTIONS"
    if name in wb.sheetnames:
        del wb[name]  # rebuild cleanly (idempotent)
    ws = wb.create_sheet(name)
    ws.append([
        "DOMAIN_SHEET_NAME", "BUSINESS_DESCRIPTION", "BUSINESS_FACTS",
        "GEOGRAPHY_LEVELS", "MANUFACTURER_SCOPE",
    ])
    for row in SHEET_DESCRIPTIONS:
        ws.append([row["domain"], row["desc"], row["facts"], row["geo"], row["mfg"]])
    print(f"  [SHEET_DESCRIPTIONS] wrote {len(SHEET_DESCRIPTIONS)} domain descriptions")


def add_geography_level(wb) -> None:
    ws = wb["DIM_GEOGRAPHY"]
    existing = {str(r[0].value).strip() for r in ws.iter_rows() if r[0].value}
    if "GEOGRAPHY_LEVEL" in existing:
        print("  [DIM_GEOGRAPHY] GEOGRAPHY_LEVEL already present")
        return
    ws.append([
        "GEOGRAPHY_LEVEL",
        "Granularity level of this geography row: Country/Region/Zone/State/"
        "Province/City/Town/Outlet.",
        "ATTR", "M", "Region>Zone>City", "Region>Province>District",
        "Region>Province>City",
    ])
    print("  [DIM_GEOGRAPHY] added GEOGRAPHY_LEVEL field")


def add_mfg_import_fields(wb) -> None:
    for sheet_name, fields in MFG_IMPORT_FIELDS.items():
        ws = wb[sheet_name]
        existing = {str(r[0].value).strip() for r in ws.iter_rows() if r[0].value}
        added = []
        for f in fields:
            if f[0] in existing:
                continue
            ws.append(list(f))
            added.append(f[0])
        print(f"  [{sheet_name}] added mfg/import fields: {added or '(none, all present)'}")


def add_long_desc_column(wb) -> None:
    """Add a BUSINESS_DESCRIPTION_LONG column to DIM_FILE and DIM_SHEET."""
    for name in ("DIM_FILE", "DIM_SHEET"):
        ws = wb[name]
        header = _header_row(ws)
        if "BUSINESS_DESCRIPTION_LONG" in header:
            print(f"  [{name}] BUSINESS_DESCRIPTION_LONG already present")
            continue
        col = ws.max_column + 1
        ws.cell(row=1, column=col, value="BUSINESS_DESCRIPTION_LONG")
        print(f"  [{name}] added BUSINESS_DESCRIPTION_LONG column at col {col}")


def main() -> None:
    wb = openpyxl.load_workbook(TPL)
    print("Updating template:", TPL)
    update_sheet_descriptions(wb)
    add_geography_level(wb)
    add_mfg_import_fields(wb)
    add_long_desc_column(wb)
    wb.save(TPL)
    print("Saved.")


if __name__ == "__main__":
    main()
