"""Parse the bundled global template workbook into a TemplateSchema.

The template (templates/global_template.xlsx) is organised as:
  - SUMMARY, PRIMARY_CONFIG, FUNCTIONAL_DOMAIN_MAP  — meta sheets
  - DIM_* sheets and domain sheets (e.g. 'PMI P&L Database') — field-definition sheets
    with header row: STANDARD_FIELD_NAME | BUSINESS_DESCRIPTION | FIELD_TYPE | FLAG
                     | INDIA_FIELD_NAME | THAILAND_FIELD_NAME | VIETNAM_FIELD_NAME

We treat every field-definition sheet as a "domain sheet" the uploaded data can
map to, and read FUNCTIONAL_DOMAIN_MAP for file→domain hints + functional area.
"""

from __future__ import annotations

from functools import lru_cache
from pathlib import Path

import openpyxl

from agent.core.settings import get_settings
from agent.models.schemas import (
    TemplateField,
    TemplateSchema,
    TemplateSheet,
)

# Sheets that are NOT field-definition domain sheets.
_META_SHEETS = {"SUMMARY", "PRIMARY_CONFIG", "FUNCTIONAL_DOMAIN_MAP", "SHEET_DESCRIPTIONS"}

# Expected header for a field-definition sheet.
_FIELD_HEADER = "STANDARD_FIELD_NAME"

_VALID_FLAGS = {"M", "C", "O"}
_VALID_TYPES = {"KEY", "ATTR", "MEAS"}


def _norm(value: object) -> str:
    return "" if value is None else str(value).strip()


def _parse_functional_domain_map(ws) -> dict[str, dict[str, str]]:
    """Return {domain_sheet_name: {COUNTRY: filename, ...}}."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}

    header = [_norm(c).upper() for c in rows[0]]
    try:
        name_idx = header.index("DOMAIN_SHEET_NAME")
    except ValueError:
        return {}

    # Country file columns look like "<COUNTRY>_FILES".
    country_cols: dict[str, int] = {}
    for idx, col in enumerate(header):
        if col.endswith("_FILES"):
            country_cols[col[: -len("_FILES")]] = idx

    result: dict[str, dict[str, str]] = {}
    for row in rows[1:]:
        if not row or name_idx >= len(row):
            continue
        domain = _norm(row[name_idx])
        if not domain:
            continue
        files = {
            country: _norm(row[idx])
            for country, idx in country_cols.items()
            if idx < len(row) and _norm(row[idx])
        }
        result[domain] = files
    return result


def _parse_sheet_descriptions(ws) -> dict[str, dict[str, str]]:
    """Return {domain_sheet_name: {desc, facts, geo, mfg}} from SHEET_DESCRIPTIONS.

    This is the file/sheet-level business meaning the classifier matches against.
    Header: DOMAIN_SHEET_NAME | BUSINESS_DESCRIPTION | BUSINESS_FACTS
            | GEOGRAPHY_LEVELS | MANUFACTURER_SCOPE
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    header = [_norm(c).upper() for c in rows[0]]

    def idx(name: str, default: int) -> int:
        try:
            return header.index(name)
        except ValueError:
            return default

    i_name = idx("DOMAIN_SHEET_NAME", 0)
    i_desc = idx("BUSINESS_DESCRIPTION", 1)
    i_facts = idx("BUSINESS_FACTS", 2)
    i_geo = idx("GEOGRAPHY_LEVELS", 3)
    i_mfg = idx("MANUFACTURER_SCOPE", 4)

    result: dict[str, dict[str, str]] = {}
    for row in rows[1:]:
        if not row or i_name >= len(row):
            continue
        name = _norm(row[i_name])
        if not name:
            continue
        result[name] = {
            "desc": _norm(row[i_desc]) if i_desc < len(row) else "",
            "facts": _norm(row[i_facts]) if i_facts < len(row) else "",
            "geo": _norm(row[i_geo]) if i_geo < len(row) else "",
            "mfg": _norm(row[i_mfg]) if i_mfg < len(row) else "",
        }
    return result


def _parse_field_sheet(ws) -> TemplateSheet | None:
    """Parse one field-definition sheet into a TemplateSheet, or None if not one."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None

    header = [_norm(c).upper() for c in rows[0]]
    if not header or header[0] != _FIELD_HEADER:
        return None

    # Column index lookup (positions are stable but resolve by name to be safe).
    def col(name: str, default: int) -> int:
        try:
            return header.index(name)
        except ValueError:
            return default

    i_name = col("STANDARD_FIELD_NAME", 0)
    i_desc = col("BUSINESS_DESCRIPTION", 1)
    i_type = col("FIELD_TYPE", 2)
    i_flag = col("FLAG", 3)

    # Any "<COUNTRY>_FIELD_NAME" columns become example hints.
    example_cols: dict[str, int] = {}
    for idx, h in enumerate(header):
        if h.endswith("_FIELD_NAME") and h != "STANDARD_FIELD_NAME":
            example_cols[h[: -len("_FIELD_NAME")]] = idx

    fields: list[TemplateField] = []
    for row in rows[1:]:
        if not row or i_name >= len(row):
            continue
        fname = _norm(row[i_name])
        if not fname:
            continue

        ftype = _norm(row[i_type]).upper() if i_type < len(row) else ""
        flag = _norm(row[i_flag]).upper() if i_flag < len(row) else ""
        examples = {
            country: _norm(row[idx])
            for country, idx in example_cols.items()
            if idx < len(row) and _norm(row[idx])
        }

        fields.append(
            TemplateField(
                standard_field_name=fname,
                business_description=_norm(row[i_desc]) if i_desc < len(row) else "",
                field_type=ftype if ftype in _VALID_TYPES else "UNKNOWN",
                flag=flag if flag in _VALID_FLAGS else "O",
                examples=examples,
            )
        )

    if not fields:
        return None
    is_dim = ws.title.upper().startswith("DIM_")
    return TemplateSheet(sheet_name=ws.title, is_dimension=is_dim, fields=fields)


def parse_template(path: str | Path | None = None) -> TemplateSchema:
    """Parse the global template workbook into a TemplateSchema."""
    path = Path(path) if path else get_settings().template_path
    if not path.exists():
        raise FileNotFoundError(
            f"Global template not found at {path}. "
            "Copy global_template.xlsx into nma/templates/."
        )

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    functional_domain_map: dict[str, dict[str, str]] = {}
    if "FUNCTIONAL_DOMAIN_MAP" in wb.sheetnames:
        functional_domain_map = _parse_functional_domain_map(wb["FUNCTIONAL_DOMAIN_MAP"])

    # Build a lookup of domain -> functional area from the map.
    area_lookup: dict[str, str] = {}
    if "FUNCTIONAL_DOMAIN_MAP" in wb.sheetnames:
        ws = wb["FUNCTIONAL_DOMAIN_MAP"]
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            header = [_norm(c).upper() for c in rows[0]]
            try:
                n_idx = header.index("DOMAIN_SHEET_NAME")
                a_idx = header.index("FUNCTIONAL_AREA")
                for row in rows[1:]:
                    if row and n_idx < len(row) and a_idx < len(row):
                        area_lookup[_norm(row[n_idx])] = _norm(row[a_idx])
            except ValueError:
                pass

    # File/sheet-level business descriptions (the canonical matching context).
    sheet_descriptions: dict[str, dict[str, str]] = {}
    if "SHEET_DESCRIPTIONS" in wb.sheetnames:
        sheet_descriptions = _parse_sheet_descriptions(wb["SHEET_DESCRIPTIONS"])

    domain_sheets: list[TemplateSheet] = []
    for ws in wb.worksheets:
        if ws.title in _META_SHEETS:
            continue
        sheet = _parse_field_sheet(ws)
        if sheet is not None:
            sheet.functional_area = area_lookup.get(sheet.sheet_name, "")
            meta = sheet_descriptions.get(sheet.sheet_name, {})
            sheet.business_description = meta.get("desc", "")
            sheet.business_facts = meta.get("facts", "")
            sheet.geography_levels = meta.get("geo", "")
            sheet.manufacturer_scope = meta.get("mfg", "")
            domain_sheets.append(sheet)

    wb.close()

    return TemplateSchema(
        version="global_template",
        domain_sheets=domain_sheets,
        functional_domain_map=functional_domain_map,
    )


@lru_cache(maxsize=2)
def load_template_cached(path_str: str | None = None) -> TemplateSchema:
    """Cached template load (keyed by path string)."""
    return parse_template(path_str)
